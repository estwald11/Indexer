"""Parsers for the formats an office archive is mostly made of: HTML, Word, Excel, PDF.

Each builds ``ParsedDocument.text`` by appending block text as it goes (the
shared ``_Builder``), so the span invariant holds by construction, and each
keeps what the format knows about structure: heading levels, list items, table
grids, page numbers. A parser that flattens a Word table into prose has
destroyed the structured path at the first stage.

``html``      Standard library. Headings, paragraphs, lists, tables, code.
``docx``      python-docx. Heading styles in English or Italian ("Titolo 1"),
              lists, tables in body order, core properties as metadata.
``xlsx``      openpyxl. One section per sheet, the used range as a table with
              its header row; capped, and says so when it is.
``pdf_text``  pypdf. Text per page with page provenance -- no layout model, so
              reading-order confidence is below 1. A PDF with no text layer (a
              scan) is handed to a configured ``fallback`` parser (Docling with
              OCR, say) instead of being indexed as nothing.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterator, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from html.parser import HTMLParser
from typing import Any, ClassVar

from indexer.core.document import BlockKind, ParsedDocument, SourceDocument, Table, TableCell
from indexer.core.provenance import PageRef
from indexer.core.registry import register
from indexer.core.stages import StageContext, parse_cache_scope
from indexer.impls.parse import _Builder, _decode
from indexer.plugin import StageImpl, dataclass_params

__all__ = ["DocxParser", "HtmlParser", "PdfTextParser", "XlsxParser", "render_table"]


def render_table(rows: Sequence[Sequence[str]], header_rows: int = 1) -> tuple[str, Table]:
    """A grid as a markdown pipe table plus its ``Table`` payload.

    The rendering is for text indexes; the payload is the structure. Cells are
    single-lined and pipes escaped, so the rendering is one row per line --
    which is what lets the segmenter split a large table by rows.
    """
    width = max((len(r) for r in rows), default=0)
    clean = [[_cell(c) for c in r] + [""] * (width - len(r)) for r in rows]
    lines = []
    for i, r in enumerate(clean):
        lines.append("| " + " | ".join(r) + " |")
        if i == max(header_rows, 1) - 1:
            lines.append("|" + "|".join("---" for _ in range(width)) + "|")
    table = Table(
        rows=tuple(
            tuple(TableCell(c, is_header=i < header_rows) for c in r) for i, r in enumerate(clean)
        ),
        header_rows=header_rows,
    )
    return "\n".join(lines), table


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split()).replace("|", "\\|")


# --------------------------------------------------------------------------- #
# html                                                                         #
# --------------------------------------------------------------------------- #


_SKIP_TAGS = ("script", "style", "nav", "footer", "header", "aside", "noscript")


@dataclass(frozen=True, slots=True)
class HtmlParams:
    encoding: str = "utf-8"
    #: Elements whose content is page furniture rather than document.
    skip_tags: list[str] = field(default_factory=lambda: list(_SKIP_TAGS))


@register(
    "parse",
    "html",
    version="1",
    params_model=dataclass_params(HtmlParams),
    summary="HTML with headings, paragraphs, lists, tables and code kept. Standard library.",
)
def _make_html(params: dict[str, Any], **_: Any) -> HtmlParser:
    return HtmlParser(params)


class _HtmlBlocks(HTMLParser):
    """Collects (kind, text, level, table_rows) in document order."""

    _HEADINGS: ClassVar[dict[str, int]] = {f"h{i}": i for i in range(1, 7)}
    _BLOCKS = frozenset({"p", "div", "section", "article", "blockquote", "dd", "dt", "figcaption"})

    def __init__(self, skip: set[str]) -> None:
        super().__init__(convert_charrefs=True)
        self.skip = skip
        self.skipping = 0
        self.out: list[tuple[str, str, int | None, list[list[str]] | None]] = []
        self.buf: list[str] = []
        self.kind = BlockKind.PARAGRAPH
        self.level: int | None = None
        self.title = ""
        self.in_title = False
        self.pre = 0
        self.table: list[list[str]] | None = None
        self.row: list[str] | None = None
        self.cell: list[str] | None = None
        self.header_cells = 0

    def _flush(self) -> None:
        text = "".join(self.buf)
        text = text.strip("\n") if self.pre else " ".join(text.split())
        if text:
            self.out.append((str(self.kind), text, self.level, None))
        self.buf = []
        self.kind, self.level = BlockKind.PARAGRAPH, None

    def handle_starttag(self, tag: str, attrs: Any) -> None:
        if tag in self.skip:
            self.skipping += 1
            return
        if self.skipping:
            return
        if tag == "title":
            self.in_title = True
        elif tag == "table":
            self._flush()
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.row = []
        elif tag in ("td", "th") and self.row is not None:
            self.cell = []
            if tag == "th" and self.table is not None and not self.table:
                self.header_cells += 1
        elif tag in self._HEADINGS:
            self._flush()
            self.kind, self.level = BlockKind.HEADING, self._HEADINGS[tag]
        elif tag == "li":
            self._flush()
            self.kind = BlockKind.LIST_ITEM
        elif tag == "pre":
            self._flush()
            self.kind, self.pre = BlockKind.CODE, self.pre + 1
        elif tag == "br":
            self.buf.append("\n" if self.pre else " ")
        elif tag in self._BLOCKS:
            self._flush()

    def handle_endtag(self, tag: str) -> None:
        if tag in self.skip:
            self.skipping = max(0, self.skipping - 1)
            return
        if self.skipping:
            return
        if tag == "title":
            self.in_title = False
        elif tag in ("td", "th") and self.cell is not None and self.row is not None:
            self.row.append(" ".join("".join(self.cell).split()))
            self.cell = None
        elif tag == "tr" and self.row is not None and self.table is not None:
            if any(c for c in self.row):
                self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            if self.table:
                self.out.append((str(BlockKind.TABLE), "", None, self.table))
            self.table, self.header_cells = None, 0
        elif tag == "pre":
            self._flush()
            self.pre = max(0, self.pre - 1)
        elif tag in self._HEADINGS or tag == "li" or tag in self._BLOCKS:
            self._flush()

    def handle_data(self, data: str) -> None:
        if self.skipping:
            return
        if self.in_title:
            self.title += data
        elif self.cell is not None:
            self.cell.append(data)
        elif self.table is None:
            self.buf.append(data)

    def close(self) -> None:
        super().close()
        self._flush()


class HtmlParser(StageImpl):
    STAGE, IMPL, VERSION = "parse", "html", "1"

    def can_parse(self, doc: SourceDocument) -> float:
        return 0.9 if doc.media_type in ("text/html", "application/xhtml+xml") else 0.1

    def parse(self, doc: SourceDocument, ctx: StageContext) -> ParsedDocument:
        html = _decode(doc, self.param("encoding", "utf-8"))
        p = _HtmlBlocks({t.lower() for t in self.param("skip_tags", _SKIP_TAGS)})
        p.feed(html)
        p.close()
        b = _Builder(doc.document_id, doc.source_uri)
        for kind, text, level, rows in p.out:
            if rows is not None:
                header = 1 if p.header_cells or len(rows) > 1 else 0
                rendered, table = render_table(rows, header_rows=header)
                b.add(rendered, BlockKind.TABLE, table=table)
            else:
                b.add(text, kind, level=level)
        meta = dict(doc.metadata)
        if p.title.strip():
            meta["title"] = " ".join(p.title.split())
        return b.finish(doc.content_hash, metadata=meta)


# --------------------------------------------------------------------------- #
# docx                                                                         #
# --------------------------------------------------------------------------- #

_HEADING_STYLE = re.compile(r"^(?:heading|titolo|intestazione|überschrift|titre)\s*(\d)", re.I)
_TITLE_STYLE = re.compile(r"^(?:title|titolo)$", re.I)
_LIST_STYLE = re.compile(r"^(?:list|elenco)", re.I)


@dataclass(frozen=True, slots=True)
class DocxParams:
    include_tables: bool = True


@register(
    "parse",
    "docx",
    version="1",
    params_model=dataclass_params(DocxParams),
    summary="Word documents: heading styles (English and Italian), lists, tables in order.",
    requires=("python-docx",),
)
def _make_docx(params: dict[str, Any], **_: Any) -> DocxParser:
    return DocxParser(params)


class DocxParser(StageImpl):
    STAGE, IMPL, VERSION = "parse", "docx", "1"

    def can_parse(self, doc: SourceDocument) -> float:
        return 0.95 if doc.media_type.endswith("wordprocessingml.document") else 0.0

    def parse(self, doc: SourceDocument, ctx: StageContext) -> ParsedDocument:
        try:
            import docx
            from docx.table import Table as DocxTable
            from docx.text.paragraph import Paragraph
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError(
                "the docx parser needs python-docx: pip install python-docx"
            ) from exc

        d = docx.Document(io.BytesIO(doc.load()))
        b = _Builder(doc.document_id, doc.source_uri)
        # Body order, not "all paragraphs then all tables": a table belongs
        # between the paragraphs that introduce and discuss it.
        for child in d.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                para = Paragraph(child, d)
                text = para.text.strip()
                if not text:
                    continue
                kind, level = _docx_kind(para)
                b.add(text, kind, level=level)
            elif tag == "tbl" and self.param("include_tables", True):
                rows = [[c.text for c in row.cells] for row in DocxTable(child, d).rows]
                rows = [_dedupe_merged(r) for r in rows if any(x.strip() for x in r)]
                if rows:
                    rendered, table = render_table(rows)
                    b.add(rendered, BlockKind.TABLE, table=table)
        meta = dict(doc.metadata)
        props = d.core_properties
        for key, value in (
            ("title", props.title),
            ("author", props.author),
            ("created", props.created),
            ("modified", props.modified),
        ):
            if value:
                meta[key] = value.date() if isinstance(value, datetime) else value
        return b.finish(doc.content_hash, metadata=meta)


def _docx_kind(para: Any) -> tuple[BlockKind, int | None]:
    style = (para.style.name if para.style is not None else "") or ""
    if m := _HEADING_STYLE.match(style):
        return BlockKind.HEADING, int(m.group(1))
    if _TITLE_STYLE.match(style):
        return BlockKind.HEADING, 1
    ppr = para._p.pPr
    if ppr is not None:
        outline = ppr.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl"
        )
        if outline is not None:
            val = outline.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
            if val is not None and val.isdigit() and int(val) < 9:
                return BlockKind.HEADING, int(val) + 1
        if ppr.numPr is not None:
            return BlockKind.LIST_ITEM, None
    if _LIST_STYLE.match(style):
        return BlockKind.LIST_ITEM, None
    return BlockKind.PARAGRAPH, None


def _dedupe_merged(row: list[str]) -> list[str]:
    """python-docx repeats a horizontally merged cell's text in every column it
    spans; keep the first and blank the repeats, as the grid shows it."""
    out: list[str] = []
    for i, c in enumerate(row):
        out.append("" if i and c == row[i - 1] and c else c)
    return out


# --------------------------------------------------------------------------- #
# xlsx                                                                         #
# --------------------------------------------------------------------------- #


@dataclass(frozen=True, slots=True)
class XlsxParams:
    #: Rows read per sheet. A ledger export can hold a million rows; indexing
    #: them as prose serves nobody, and the truncation is recorded.
    max_rows: int = 5000
    max_columns: int = 60
    include_hidden_sheets: bool = False


@register(
    "parse",
    "xlsx",
    version="1",
    params_model=dataclass_params(XlsxParams),
    summary="Excel workbooks: a section per sheet, the used range as a table with its header.",
    requires=("openpyxl",),
)
def _make_xlsx(params: dict[str, Any], **_: Any) -> XlsxParser:
    return XlsxParser(params)


class XlsxParser(StageImpl):
    STAGE, IMPL, VERSION = "parse", "xlsx", "1"

    def can_parse(self, doc: SourceDocument) -> float:
        return 0.95 if doc.media_type.endswith("spreadsheetml.sheet") else 0.0

    def parse(self, doc: SourceDocument, ctx: StageContext) -> ParsedDocument:
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("the xlsx parser needs openpyxl: pip install openpyxl") from exc

        wb = openpyxl.load_workbook(io.BytesIO(doc.load()), read_only=True, data_only=True)
        max_rows = int(self.param("max_rows", 5000))
        max_cols = int(self.param("max_columns", 60))
        b = _Builder(doc.document_id, doc.source_uri)
        truncated: list[str] = []
        for ws in wb.worksheets:
            if ws.sheet_state != "visible" and not self.param("include_hidden_sheets", False):
                continue
            rows: list[list[str]] = []
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= max_rows:
                    truncated.append(ws.title)
                    break
                cells = [_cell(v) for v in row[:max_cols]]
                if any(cells):
                    rows.append(cells)
            if not rows:
                continue
            while rows and all(not r[-1] for r in rows):  # trailing empty columns
                rows = [r[:-1] for r in rows]
            b.add(ws.title, BlockKind.HEADING, level=1)
            rendered, table = render_table(rows)
            b.add(rendered, BlockKind.TABLE, table=table)
        wb.close()
        meta = dict(doc.metadata)
        if truncated:
            meta["truncated_sheets"] = truncated
        return b.finish(doc.content_hash, metadata=meta)


# --------------------------------------------------------------------------- #
# pdf                                                                          #
# --------------------------------------------------------------------------- #


@dataclass(frozen=True, slots=True)
class PdfTextParams:
    #: Below this many characters per page on average, the PDF is treated as
    #: having no text layer -- a scan -- and handed to ``fallback``.
    min_chars_per_page: int = 20
    #: A parser spec ({impl, params}) for PDFs without a text layer. Without
    #: one, a scan is indexed as its metadata alone and flagged.
    fallback: dict[str, Any] | None = None
    #: Keep the text layer's line breaks inside a paragraph instead of joining
    #: its lines with spaces. A segmenter that reads line starts -- the item
    #: codes of a specification (``items``) -- needs them; search reads the same.
    keep_line_breaks: bool = False


@register(
    "parse",
    "pdf_text",
    version="1",
    params_model=dataclass_params(PdfTextParams),
    summary=(
        "Text-layer PDFs via pypdf, page by page with page provenance. Scans go to a "
        "configured fallback (OCR) instead of being indexed as nothing."
    ),
    requires=("pypdf",),
)
def _make_pdf_text(params: dict[str, Any], **_: Any) -> PdfTextParser:
    return PdfTextParser(params)


class PdfTextParser(StageImpl):
    STAGE, IMPL, VERSION = "parse", "pdf_text", "1"

    def __init__(self, params: dict[str, Any]) -> None:
        super().__init__(params)
        self._fallback: Any = None
        spec = params.get("fallback")
        if spec:
            from indexer.config.loader import resolve_impl

            reg, norm = resolve_impl(
                "parse", {"impl": spec["impl"], "params": spec.get("params", {})}
            )
            self._fallback = reg.build(norm)

    def fingerprint(self) -> Any:
        from indexer.core.accounting import StageFingerprint
        from indexer.core.ids import hash_obj

        # The fallback parses some of this parser's documents, so its version
        # is part of what this parser produces.
        return StageFingerprint(
            stage=self.STAGE,
            impl=self.IMPL,
            version=self.VERSION,
            params_hash=hash_obj(
                {
                    "params": self._params,
                    "fallback": self._fallback.fingerprint().key() if self._fallback else None,
                }
            ),
        )

    def cache_scope(self, doc: SourceDocument) -> str:
        fb = parse_cache_scope(self._fallback, doc) if self._fallback else ""
        return f"{doc.media_type}|{fb}"

    def can_parse(self, doc: SourceDocument) -> float:
        return 0.8 if doc.media_type == "application/pdf" else 0.0

    def parse(self, doc: SourceDocument, ctx: StageContext) -> ParsedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("the pdf_text parser needs pypdf: pip install pypdf") from exc

        reader = PdfReader(io.BytesIO(doc.load()))
        pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
        chars = sum(len(t.strip()) for _, t in pages)
        min_chars = int(self.param("min_chars_per_page", 20))
        if pages and chars < min_chars * len(pages) and self._fallback is not None:
            parsed: ParsedDocument = self._fallback.parse(doc, ctx)
            return parsed
        b = _Builder(doc.document_id, doc.source_uri)
        keep = bool(self.param("keep_line_breaks", False))
        for page_no, text in pages:
            for para in _paragraphs(text, keep_line_breaks=keep):
                block = b.add(para, BlockKind.PARAGRAPH)
                b.blocks[-1] = _with_page(block, page_no)
        meta = dict(doc.metadata)
        meta["text_layer"] = chars >= min_chars * max(1, len(pages))
        info = reader.metadata
        if info is not None and info.title:
            meta["title"] = str(info.title)
        return b.finish(
            doc.content_hash,
            page_count=len(pages),
            # Text-layer order without a layout model: right for single-column
            # documents, wrong for many forms and multi-column layouts. Said so.
            reading_order_confidence=0.7,
            metadata=meta,
        )


def _paragraphs(text: str, *, keep_line_breaks: bool = False) -> Iterator[str]:
    """Paragraphs from a page's extracted text: blank lines separate them, and
    lines hyphenated at the margin are rejoined. A paragraph's lines are joined
    with spaces, or kept on their own lines with ``keep_line_breaks``."""
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    sep = "\n" if keep_line_breaks else " "
    for para in re.split(r"\n\s*\n", text):
        joined = sep.join(line.strip() for line in para.splitlines() if line.strip())
        if joined:
            yield joined


def _with_page(block: Any, page: int) -> Any:
    from dataclasses import replace

    return replace(block, provenance=replace(block.provenance, pages=PageRef.of(page)))
