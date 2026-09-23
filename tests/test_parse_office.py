"""Office-format parsers keep structure and satisfy the parse contract.

Every document is generated here, and every parse is checked with
``check_parsed_document`` -- spans resolve, tables carry grids, headings carry
levels -- because a parser that returns plausible text with broken spans is
the failure that surfaces three stages later as a wrong citation.
"""

from __future__ import annotations

import io
from typing import Any

import pytest

from indexer.core.accounting import InMemoryAccountant
from indexer.core.cache import NullCache
from indexer.core.document import BlockKind, SourceDocument
from indexer.core.ids import DocumentId, hash_bytes
from indexer.core.stages import StageContext
from indexer.eval.checks import check_parsed_document
from indexer.impls.containers import media_type_for
from indexer.impls.parse_office import DocxParser, HtmlParser, PdfTextParser, XlsxParser

CTX = StageContext(cache=NullCache(), accountant=InMemoryAccountant())


class _Bytes(SourceDocument):
    __slots__ = ()

    def load(self) -> bytes:
        return self.metadata["_bytes"]  # type: ignore[no-any-return]


def _doc(name: str, data: bytes) -> SourceDocument:
    return _Bytes(
        document_id=DocumentId(name),
        source_uri=f"mem://{name}",
        content_hash=hash_bytes(data),
        media_type=media_type_for(name),
        size_bytes=len(data),
        metadata={"_bytes": data, "name": name},
    )


def _parse(parser: Any, name: str, data: bytes) -> Any:
    parsed = parser.parse(_doc(name, data), CTX)
    assert check_parsed_document(parsed) == [], name
    return parsed


def _kinds(parsed: Any) -> list[tuple[str, str]]:
    return [(str(b.kind), b.text) for b in parsed.blocks]


class TestHtml:
    def test_structure_is_kept(self) -> None:
        html = """<html><head><title>Listino 2025</title><style>p{}</style></head><body>
        <nav>Home | Chi siamo</nav>
        <h1>Listino prezzi</h1><p>Prezzi validi dal <b>1° gennaio</b> 2025.</p>
        <h2>Servizi</h2><ul><li>Consulenza</li><li>Formazione</li></ul>
        <table><tr><th>Servizio</th><th>Prezzo</th></tr>
               <tr><td>Consulenza</td><td>1.200,00 €</td></tr></table>
        <pre>codice  esempio</pre><footer>© Azienda</footer></body></html>"""
        parsed = _parse(HtmlParser({}), "listino.html", html.encode())
        kinds = _kinds(parsed)
        assert kinds[0] == ("heading", "Listino prezzi")
        assert ("paragraph", "Prezzi validi dal 1° gennaio 2025.") in kinds
        assert ("list_item", "Consulenza") in kinds
        assert parsed.metadata["title"] == "Listino 2025"
        table = next(b for b in parsed.blocks if str(b.kind) == BlockKind.TABLE)
        assert table.table is not None and table.table.shape == (2, 2)
        assert table.table.rows[1][1].text == "1.200,00 €"
        assert ("code", "codice  esempio") in kinds
        # Page furniture is not document.
        assert not any("Chi siamo" in t or "©" in t for _, t in kinds)


class TestDocx:
    def test_headings_lists_and_tables_in_body_order(self) -> None:
        docx = pytest.importorskip("docx")
        from docx.enum.style import WD_STYLE_TYPE

        d = docx.Document()
        d.core_properties.title = "Contratto quadro"
        d.core_properties.author = "Ufficio legale"
        d.add_heading("Contratto quadro", level=1)
        d.add_paragraph("Tra le parti si conviene quanto segue.")
        italian = d.styles.add_style("Titolo 2", WD_STYLE_TYPE.PARAGRAPH)
        d.add_paragraph("Oggetto", style=italian)
        d.add_paragraph("Fornitura di servizi", style="List Bullet")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Voce", "Importo"
        t.cell(1, 0).text, t.cell(1, 1).text = "Canone", "12.000,00"
        d.add_paragraph("Il presente contratto decorre dalla firma.")
        buf = io.BytesIO()
        d.save(buf)

        parsed = _parse(DocxParser({}), "contratto.docx", buf.getvalue())
        kinds = [str(b.kind) for b in parsed.blocks]
        assert kinds == ["heading", "paragraph", "heading", "list_item", "table", "paragraph"]
        levels = [b.level for b in parsed.blocks if str(b.kind) == "heading"]
        assert levels == [1, 2]  # "Titolo 2", the Italian style name, is a heading
        table = parsed.blocks[4].table
        assert table is not None and table.rows[1][1].text == "12.000,00"
        assert parsed.metadata["title"] == "Contratto quadro"
        assert parsed.metadata["author"] == "Ufficio legale"


class TestXlsx:
    def test_one_section_per_sheet_with_a_header_row(self) -> None:
        openpyxl = pytest.importorskip("openpyxl")
        from datetime import date

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Fatture"
        ws.append(["Numero", "Data", "Importo"])
        ws.append(["2025/1", date(2025, 1, 15), 1250.5])
        ws.append(["2025/2", date(2025, 2, 3), 99.0])
        hidden = wb.create_sheet("Appunti")
        hidden.append(["non indicizzare"])
        hidden.sheet_state = "hidden"
        buf = io.BytesIO()
        wb.save(buf)

        parsed = _parse(XlsxParser({}), "fatture.xlsx", buf.getvalue())
        assert _kinds(parsed)[0] == ("heading", "Fatture")
        table = parsed.blocks[1].table
        assert table is not None and table.header_rows == 1
        assert [c.text for c in table.rows[1]] == ["2025/1", "2025-01-15", "1250.5"]
        assert [c.text for c in table.rows[2]] == ["2025/2", "2025-02-03", "99"]
        assert not any("indicizzare" in b.text for b in parsed.blocks)

    def test_truncation_is_recorded(self) -> None:
        openpyxl = pytest.importorskip("openpyxl")
        wb = openpyxl.Workbook()
        for i in range(30):
            wb.active.append([f"riga {i}", i])
        buf = io.BytesIO()
        wb.save(buf)
        parsed = _parse(XlsxParser({"max_rows": 10}), "big.xlsx", buf.getvalue())
        assert parsed.metadata["truncated_sheets"] == ["Sheet"]


def _pdf(pages: list[str]) -> bytes:
    """A minimal PDF, one text line per page, Helvetica, so text extracts."""
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # pages, filled in below
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
        content_ref = len(objects)
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_ref
        )
        kids.append(len(objects))
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % k for k in kids),
        len(kids),
    )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (i, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


class TestPdf:
    def test_text_pages_carry_page_numbers(self) -> None:
        pytest.importorskip("pypdf")
        data = _pdf(["Prima pagina del verbale di assemblea", "Seconda pagina con le delibere"])
        parsed = _parse(PdfTextParser({}), "verbale.pdf", data)
        assert parsed.page_count == 2
        assert parsed.metadata["text_layer"] is True
        pages = [(b.provenance.pages.start, b.text) for b in parsed.blocks if b.provenance.pages]
        assert pages == [
            (1, "Prima pagina del verbale di assemblea"),
            (2, "Seconda pagina con le delibere"),
        ]
        assert parsed.reading_order_confidence < 1.0

    def test_a_scan_goes_to_the_fallback(self) -> None:
        pytest.importorskip("pypdf")
        scan = _pdf(["", ""])  # pages with no text layer
        parser = PdfTextParser({"fallback": {"impl": "passthrough"}})
        parsed = parser.parse(_doc("scansione.pdf", scan), CTX)
        # The passthrough fallback decodes bytes: what matters is that the
        # fallback, not the text extractor, produced the parse.
        assert len(parsed.blocks) == 1 and str(parsed.blocks[0].kind) == "other"
        assert parser.fingerprint() != PdfTextParser({}).fingerprint()
